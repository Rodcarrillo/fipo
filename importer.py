"""Lectura de estados de cuenta (CSV, Excel, PDF) -> filas normalizadas para revisión.

Nada se guarda aquí: el resultado siempre pasa por una pantalla de revisión y el usuario confirma.
Los importes salen en centavos con signo (positivo = entra dinero, negativo = sale).
"""
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

MAX_ROWS = 5000
MAX_PDF_PAGES = 60
MONTHS = {"ene": 1, "jan": 1, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8, "aug": 8,
          "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12, "dec": 12}
HEADER_WORDS = {
    "date": ("fecha", "date", "f. operacion", "f.operacion", "fecha operacion", "fecha de operacion", "fecha valor", "posting"),
    "description": ("descripcion", "descripción", "concepto", "detalle", "description", "referencia", "movimiento", "memo", "details"),
    "amount": ("importe", "monto", "amount", "cantidad", "valor"),
    "debit": ("cargo", "cargos", "retiro", "retiros", "debito", "débito", "debit", "withdrawal", "egreso"),
    "credit": ("abono", "abonos", "deposito", "depósito", "depositos", "depósitos", "credito", "crédito", "credit", "deposit", "ingreso"),
    "balance": ("saldo", "balance"),
}
INCOME_WORDS = ("abono", "deposito", "depósito", "nomina", "nómina", "pago recibido", "recibido", "interes ganado", "interés ganado",
                "reembolso", "devolucion", "devolución", "cashback", "dividend", "payroll", "refund", "salary")
OPENING_RE = re.compile(r"(saldo\s+(anterior|inicial)|balance\s+(forward|anterior)|opening\s+balance|beginning\s+balance)", re.I)


class ImportError_(Exception):
    """Error legible para el usuario."""


# ------------------------------------------------------------------ números y fechas
def parse_amount(raw):
    """Devuelve centavos con signo o None. Acepta $1,234.56 · 1.234,56 · (45.00) · 45.00- · 45.00 CR/DB."""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float, Decimal)):
        try:
            return int((Decimal(str(raw)) * 100).quantize(Decimal(1)))
        except InvalidOperation:
            return None
    s = str(raw).strip()
    if not s:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    up = s.upper()
    if up.endswith("DB") or up.endswith("D"):
        if re.search(r"\d\s*(DB|D)$", up):
            neg, s = True, re.sub(r"\s*(DB|D)$", "", s, flags=re.I)
    elif re.search(r"\d\s*(CR|C)$", up):
        s = re.sub(r"\s*(CR|C)$", "", s, flags=re.I)
    s = re.sub(r"[A-Za-z$€£\s]", "", s)
    if s.endswith("-"):
        neg, s = True, s[:-1]
    if s.startswith("-"):
        neg, s = True, s[1:]
    elif s.startswith("+"):
        s = s[1:]
    if not s or not re.fullmatch(r"[\d.,]+", s):
        return None
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        thou = "." if dec == "," else ","
        s = s.replace(thou, "").replace(dec, ".")
    elif "," in s:
        if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")
        elif re.fullmatch(r"\d+,\d{1,2}", s):
            s = s.replace(",", ".")
        else:
            return None
    elif s.count(".") > 1:
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        else:
            return None
    try:
        cents = int((Decimal(s) * 100).quantize(Decimal(1)))
    except InvalidOperation:
        return None
    return -cents if neg else cents


def _mk_date(y, m, d):
    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_date_value(raw, day_first=True, default_year=None):
    """Devuelve 'AAAA-MM-DD' o None."""
    if raw is None:
        return None
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        if 20000 < raw < 80000:                      # serie de Excel
            return (date(1899, 12, 30) + timedelta(days=int(raw))).isoformat()
        return None
    s = str(raw).strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        d = _mk_date(int(m[1]), int(m[2]), int(m[3]))
        return d.isoformat() if d else None
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), int(m[3])
        y = y + 2000 if y < 100 else y
        day, mon = (a, b) if day_first else (b, a)
        if mon > 12 and day <= 12:                   # el orden elegido no cuadra: se intenta el otro
            day, mon = mon, day
        d = _mk_date(y, mon, day)
        return d.isoformat() if d else None
    m = re.match(r"^(\d{1,2})[\s/.-]*([A-Za-záéíóúÁÉÍÓÚ]{3,})\.?[\s/.,-]*(\d{2,4})?", s)
    if m and m[2][:3].lower() in MONTHS:
        y = int(m[3]) if m[3] else default_year
        if y is None:
            return None
        y = y + 2000 if y < 100 else y
        d = _mk_date(y, MONTHS[m[2][:3].lower()], int(m[1]))
        return d.isoformat() if d else None
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})$", s)   # sin año (PDF)
    if m and default_year:
        a, b = int(m[1]), int(m[2])
        day, mon = (a, b) if day_first else (b, a)
        d = _mk_date(default_year, mon, day)
        return d.isoformat() if d else None
    return None


def clean_desc(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()[:200]


def norm_desc(v):
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


# ------------------------------------------------------------------ tablas (CSV / Excel)
def _detect_columns(header):
    found = {}
    for idx, name in enumerate(header):
        n = str(name or "").strip().lower()
        if not n:
            continue
        for key, words in HEADER_WORDS.items():
            if key in found:
                continue
            if any(n == w or n.startswith(w) or w in n for w in words):
                found[key] = idx
                break
    return found


def _find_header(grid):
    best, best_i = None, None
    for i, row in enumerate(grid[:40]):
        cols = _detect_columns(row)
        score = len(cols) + (2 if "date" in cols else 0)
        if "date" in cols and ("amount" in cols or "debit" in cols or "credit" in cols):
            if best is None or score > best[0]:
                best, best_i = (score, cols), i
    return (best_i, best[1]) if best else (None, {})


def _rows_from_grid(grid, mapping, header_index, day_first, default_year, invert):
    rows, skipped = [], 0
    start = 0 if header_index is None else header_index + 1
    for raw in grid[start:start + MAX_ROWS]:
        get = lambda k: raw[mapping[k]] if k in mapping and mapping[k] < len(raw) else None
        d = parse_date_value(get("date"), day_first, default_year)
        amt = None
        if "amount" in mapping:
            amt = parse_amount(get("amount"))
        if amt is None and ("debit" in mapping or "credit" in mapping):
            deb, cre = parse_amount(get("debit")), parse_amount(get("credit"))
            if deb not in (None, 0) or cre not in (None, 0):
                amt = (abs(cre) if cre else 0) - (abs(deb) if deb else 0)
        if d is None or amt is None or amt == 0:
            if any(str(c or "").strip() for c in raw):
                skipped += 1
            continue
        if invert:
            amt = -amt
        rows.append({"date": d, "description": clean_desc(get("description")), "amount": amt,
                     "balance": parse_amount(get("balance")) if "balance" in mapping else None, "confidence": "column"})
    return rows, skipped


def _decode(data):
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("latin-1", errors="replace")


def _grid_csv(data):
    text = _decode(data)
    sample = text[:8000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delim = max(",;\t|", key=lambda d: sample.count(d))
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _grid_excel(data):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError_("Falta instalar 'openpyxl' en el servidor para leer Excel (pip install openpyxl).")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportError_("No se pudo abrir el archivo de Excel. Verifica que sea .xlsx y no esté protegido con contraseña.")
    for ws in wb.worksheets:
        grid = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS + 60:
                break
            grid.append(list(row))
        if any(any(c not in (None, "") for c in r) for r in grid):
            wb.close()
            return grid
    wb.close()
    raise ImportError_("El archivo de Excel no tiene datos.")


def parse_table(grid, mapping=None, day_first=True, default_year=None, invert=False):
    width = max((len(r) for r in grid), default=0)
    hi, auto = _find_header(grid)
    used = mapping if mapping else auto
    header = grid[hi] if hi is not None else (grid[0] if grid else [])
    columns = [clean_desc(header[i]) if i < len(header) and clean_desc(header[i]) else f"Columna {i + 1}" for i in range(width)]
    out = {"columns": columns, "mapping": used, "header_row": hi, "warnings": []}
    if not mapping and not ("date" in used and ("amount" in used or "debit" in used or "credit" in used)):
        out.update(rows=[], needs_mapping=True)
        out["sample"] = [[clean_desc(c) for c in r] for r in grid[:8]]
        return out
    if mapping:
        hi = hi if hi is not None else 0
    rows, skipped = _rows_from_grid(grid, used, hi, day_first, default_year, invert)
    if skipped:
        out["warnings"].append(f"Se omitieron {skipped} filas sin fecha o importe válidos (encabezados, totales, saldos).")
    out.update(rows=rows, needs_mapping=False)
    out["sample"] = [[clean_desc(c) for c in r] for r in grid[:8]]
    return out


# ------------------------------------------------------------------ PDF (mejor esfuerzo)
DATE_TOKEN = r"(?:\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}|\d{4}-\d{2}-\d{2}|\d{1,2}[ /-][A-Za-z]{3}[A-Za-z]*\.?(?:[ /-]\d{2,4})?|\d{1,2}[/-]\d{1,2})"
MONEY_TOKEN = r"(?<![\w/])\(?-?\$?\s?\d[\d,]*\.\d{2}\)?-?(?:\s?(?:CR|DB))?(?![\w/])"
LINE_RE = re.compile(rf"^\s*(?P<date>{DATE_TOKEN})\s+(?P<rest>.+)$")


def parse_pdf(data, day_first=True, default_year=None, invert=False):
    try:
        import pdfplumber
    except ImportError:
        raise ImportError_("Falta instalar 'pdfplumber' en el servidor para leer PDF (pip install pdfplumber).")
    lines, pages = [], 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:MAX_PDF_PAGES]:
                pages += 1
                t = page.extract_text() or ""
                lines.extend(t.split("\n"))
    except Exception:
        raise ImportError_("No se pudo leer el PDF. Puede estar protegido con contraseña o dañado.")
    if len("".join(lines).strip()) < 40:
        raise ImportError_("Este PDF no contiene texto seleccionable (parece escaneado). Descarga el estado de cuenta en CSV o Excel desde tu banco.")
    text_all = "\n".join(lines)
    if default_year is None:
        yrs = re.findall(r"\b(20[0-3]\d)\b", text_all)
        default_year = int(max(set(yrs), key=yrs.count)) if yrs else date.today().year
    opening = None
    parsed = []
    for line in lines:
        if opening is None and OPENING_RE.search(line):
            toks = re.findall(MONEY_TOKEN, line)
            if toks:
                opening = parse_amount(toks[-1])
            continue
        m = LINE_RE.match(line)
        if not m:
            continue
        d = parse_date_value(m["date"], day_first, default_year)
        if not d:
            continue
        rest = m["rest"]
        rest = re.sub(rf"^{DATE_TOKEN}\s+", "", rest)          # segunda fecha (fecha valor)
        toks = list(re.finditer(MONEY_TOKEN, rest))
        if not toks:
            continue
        first = toks[0].start()
        desc = clean_desc(rest[:first])
        nums = [parse_amount(t.group(0)) for t in toks[-3:]]
        if any(n is None for n in nums):
            continue
        amt, bal = (nums[0], None) if len(nums) == 1 else (nums[-2], nums[-1]) if len(nums) == 2 else (nums[0] or nums[1], nums[-1])
        if amt == 0:
            continue
        parsed.append({"date": d, "description": desc, "raw_amount": amt, "balance": bal, "explicit_sign": bool(re.search(r"-|\(|DB|CR", toks[-2 if len(nums) >= 2 else -1].group(0)))})
    if not parsed:
        return {"columns": [], "mapping": {}, "rows": [], "warnings": [], "needs_mapping": False,
                "error": "No encontré movimientos en este PDF. Cada banco usa un formato distinto; prueba con el CSV o Excel de tu banco."}
    newest_first = parsed[0]["date"] > parsed[-1]["date"]
    seq = list(reversed(parsed)) if newest_first else parsed
    rows, prev = [], opening
    for p in seq:
        a = abs(p["raw_amount"])
        conf, signed = "default", -a
        if p["balance"] is not None and prev is not None and abs(abs(p["balance"] - prev) - a) <= 1:
            signed, conf = (a if p["balance"] > prev else -a), "balance"
        elif p["explicit_sign"] and p["raw_amount"] < 0:
            signed, conf = -a, "column"
        elif any(w in p["description"].lower() for w in INCOME_WORDS):
            signed, conf = a, "keyword"
        if p["balance"] is not None:
            prev = p["balance"]
        elif prev is not None:
            prev += signed
        rows.append({"date": p["date"], "description": p["description"], "amount": -signed if invert else signed, "balance": p["balance"], "confidence": conf})
    warns = ["Lectura de PDF en modo de mejor esfuerzo: revisa cada fila, sobre todo las marcadas con signo dudoso, antes de guardar."]
    if any(r["confidence"] in ("default", "keyword") for r in rows):
        warns.append("En algunas filas no pude confirmar si es ingreso o gasto; están marcadas para que las verifiques.")
    return {"columns": [], "mapping": {}, "rows": sorted(rows, key=lambda r: r["date"]), "warnings": warns, "needs_mapping": False, "pages": pages}


# ------------------------------------------------------------------ entrada principal
def parse_statement(filename, data, mapping=None, day_first=True, default_year=None, invert=False):
    name = (filename or "").lower()
    if len(data) == 0:
        raise ImportError_("El archivo está vacío.")
    if name.endswith((".csv", ".txt", ".tsv")):
        kind, res = "csv", parse_table(_grid_csv(data), mapping, day_first, default_year, invert)
    elif name.endswith((".xlsx", ".xlsm")):
        kind, res = "excel", parse_table(_grid_excel(data), mapping, day_first, default_year, invert)
    elif name.endswith(".xls"):
        raise ImportError_("El formato .xls antiguo no es compatible. Abre el archivo y guárdalo como .xlsx o .csv.")
    elif name.endswith(".pdf") or data[:5] == b"%PDF-":
        kind, res = "pdf", parse_pdf(data, day_first, default_year, invert)
    else:
        raise ImportError_("Formato no compatible. Usa CSV, Excel (.xlsx) o PDF.")
    res["kind"] = kind
    if len(res.get("rows", [])) > 2000:
        raise ImportError_("El archivo tiene más de 2,000 movimientos. Divídelo por periodos e impórtalos por partes.")
    return res
